import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_full_project_report():
    doc = Document()

    # 1. Page Setup (Normal 1-inch margins)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # XML Helper for cell shading
    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    # XML Helper for cell margins
    def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="{top}" w:type="dxa"/>
                <w:bottom w:w="{bottom}" w:type="dxa"/>
                <w:left w:w="{left}" w:type="dxa"/>
                <w:right w:w="{right}" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(tcMar)

    # 2. Cover / Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(8)
    title_run = title_p.add_run("INTEGRATED BUSINESS DECISION SUPPORT SYSTEM (DSS)")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(16, 44, 87) # Deep Navy

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(24)
    sub_run = sub_p.add_run("A Unified Commercial Management Platform for Tyre Sales, EV Charging & Fleet Operations\nCase Study: Multi-Sector Commercial Hub in Piliyandala, Sri Lanka")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(80, 80, 80)

    # Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Project Status", "Fully Implemented & Verified MVP (Phase 1 & 2 Complete)"),
        ("Architecture", "Three-Tier RESTful Client-Server with Standalone Analytics"),
        ("Deployment Profile", "Localhost (Node.js Port 5000 / React Port 5173 / MySQL 3306)"),
        ("Documentation Scope", "Architecture, Security, Database Schemas, API Specs & DSS Logic")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width = Inches(2.2)
        c2.width = Inches(4.3)
        set_cell_background(c1, "F0F4F8")
        set_cell_background(c2, "FAFAFA")
        set_cell_margins(c1, top=100, bottom=100, left=120, right=120)
        set_cell_margins(c2, top=100, bottom=100, left=120, right=120)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(k)
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(v)
        r2.font.size = Pt(9.5)

    doc.add_page_break()

    # Style Helpers
    def add_sec_heading(text, level=1):
        h = doc.add_heading(text, level=level)
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        for r in h.runs:
            r.font.name = "Calibri"
            r.font.color.rgb = RGBColor(16, 44, 87) if level == 1 else RGBColor(40, 70, 120)
        return h

    def add_body_p(text, bold_prefix=None, bullet=False):
        style = 'List Bullet' if bullet else 'Normal'
        p = doc.add_paragraph(style=style)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Calibri"
            r_bold.font.size = Pt(10.5)
            r_bold.font.bold = True
        r_text = p.add_run(text)
        r_text.font.name = "Calibri"
        r_text.font.size = Pt(10.5)
        return p

    def add_callout(text, prefix="OPERATIONAL DIRECTIVE: "):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        cell.width = Inches(6.5)
        set_cell_background(cell, "EBF3FB")
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r1 = p.add_run(prefix)
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(16, 44, 87)
        r2 = p.add_run(text)
        r2.font.size = Pt(10)

    # -------------------------------------------------------------
    # SECTION 1: EXECUTIVE SUMMARY & PROBLEM STATEMENT
    # -------------------------------------------------------------
    add_sec_heading("1. Executive Summary & Problem Context")
    add_body_p(
        "Commercial enterprises operating across high-capital, disparate automotive sub-sectors face significant management friction when each operation functions in an administrative silo. In Piliyandala, a key suburban transit and commercial node in the Western Province of Sri Lanka, commercial operators commonly administer three interconnected vehicle services: (1) an automotive tyre retail, repair, and wheel-alignment workshop; (2) an electric vehicle (EV) commercial charging point; and (3) a private fleet operating on commercial ride-hailing networks (Uber and PickMe)."
    )
    add_body_p(
        "When managed using conventional fragmented paper registers or separate isolated POS and accounting utilities, executives face four systemic operational vulnerabilities:"
    )
    add_body_p("Cash Flow Obfuscation: Net profits from high-margin retail tyre sales mask heavy operational fuel burn and maintenance overheads in the ride-hailing fleet.", bold_prefix="1. ", bullet=False)
    add_body_p("Inventory Vulnerability: Tyres sold via cash walk-ins without real-time row-locking updates lead to stockouts on fast-moving commuter tire sizes (e.g., 195/65R15, 205/55R16).", bold_prefix="2. ", bullet=False)
    add_body_p("Unmonitored Energy Tariffs: Grid consumption tariffs for commercial EV charging fluctuated without immediate linkage to daily transaction receipts.", bold_prefix="3. ", bullet=False)
    add_body_p("Lack of Synergy Intelligence: Fleet owners fail to quantitatively evaluate how rapidly converting their internal combustion engine (ICE) fleet to electric power would save money compared to public EV charging revenue.", bold_prefix="4. ", bullet=False)
    add_body_p(
        "This project resolves these challenges by constructing a full-stack Integrated Business Decision Support System (DSS) that centralizes operational transaction pipelines, maintains automated stock decrements, enforces Role-Based Access Control (RBAC), and computes real-time cross-sector synergy heuristics."
    )

    # -------------------------------------------------------------
    # SECTION 2: SYSTEM ARCHITECTURE & TECHNOLOGY STACK
    # -------------------------------------------------------------
    add_sec_heading("2. System Architecture & Technical Specifications")
    add_body_p(
        "The application is engineered on an asynchronous, modular three-tier client-server architectural model, augmented by an offline analytical audit pipeline."
    )

    tech_table = doc.add_table(rows=5, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Architectural Layer", "Technology Selection", "Technical Responsibility & Justification"]
    for j, h in enumerate(headers):
        cell = tech_table.rows[0].cells[j]
        set_cell_background(cell, "102C57")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    tech_rows = [
        ("Presentation Layer (Client)", "React 18, Vite, Bootstrap 5, Chart.js", "Provides a reactive, low-latency Single Page Application (SPA). Employs role-aware tab navigation and reactive chart visualization."),
        ("Application Layer (REST API)", "Node.js, Express.js, JWT, Bcrypt.js", "Handles stateless RESTful requests, input validation, bcrypt password hashing, and ACID-compliant multi-table database transactions."),
        ("Data Persistence Layer", "MySQL 8.0, InnoDB Storage Engine", "Maintains relational integrity, foreign key cascades, check constraints, and row-level locks (FOR UPDATE) to prevent concurrency anomalies."),
        ("Audit & Data Science Engine", "Python 3, Pandas, SQLAlchemy, PyMySQL", "Performs direct analytical extraction, calculates statistical cross-sector correlation metrics, and exports datestamped CSV audit trails.")
    ]
    for i, row in enumerate(tech_rows):
        r_cells = tech_table.rows[i+1].cells
        for j, val in enumerate(row):
            cell = r_cells[j]
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            set_cell_background(cell, "F8F9FA" if i % 2 == 0 else "FFFFFF")
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if j == 0:
                r.font.bold = True

    add_sec_heading("2.1 High-Level Data Flow Architecture", level=2)
    add_body_p(
        "[Client: React SPA] ---> HTTP/JSON Requests with Bearer JWT ---> [Gateway: Express.js (Port 5000)]\n"
        "                                                                       |\n"
        "                     +-------------------+-----------------------------+\n"
        "                     |                   |                             |\n"
        "             [Route: /api/auth]   [Route: /api/tyres]        [Route: /api/analytics]\n"
        "                     |                   |                             |\n"
        "             (Bcrypt Verification) (ACID Transactions)        (Cross-Sector Aggregations)\n"
        "                     +-------------------+-----------------------------+\n"
        "                                         v\n"
        "                        [Database Tier: MySQL (InnoDB)]\n"
        "                                         ^\n"
        "                                         |\n"
        "             [Analytics Tier: Python Pandas Script via SQLAlchemy Engine]"
    )

    # -------------------------------------------------------------
    # SECTION 3: FUNCTIONAL MODULES & HOW THE SYSTEM WORKS
    # -------------------------------------------------------------
    add_sec_heading("3. Functional Modules & Operational Workflows")

    add_sec_heading("3.1 Role-Based Access Control (RBAC) & Authentication", level=2)
    add_body_p(
        "To ensure business confidentiality and enforce operational separation of duties, the system implements tokenized JWT authentication across four standard organizational roles:"
    )
    add_body_p("Business Owner & Executive Accountant: Unrestricted authorization. Access to Executive Dashboards, sector comparisons, financial print utilities, and configuration controls.", bold_prefix="• Owner: ", bullet=True)
    add_body_p("Retail Workshop Supervisor: Restricted strictly to Tyre Inventory additions, stock valuation, and POS checkout workflows.", bold_prefix="• Shop Manager: ", bullet=True)
    add_body_p("Charging Station Attendant: Dedicated access to charging point status toggles and real-time kWh billing inputs.", bold_prefix="• EV Operator: ", bullet=True)
    add_body_p("Assigned Vehicle Operators: Dedicated view for submitting trip gross earnings (Uber/PickMe receipts) and recording fuel/maintenance expense vouchers.", bold_prefix="• Driver / Fleet Manager: ", bullet=True)

    add_sec_heading("3.2 Executive Decision Support Dashboard & Strategic Rules", level=2)
    add_body_p(
        "The executive view aggregates all operational metrics into a centralized real-time dashboard. In addition to financial summaries, the system executes an automated rule-based heuristic inference engine:"
    )
    add_callout(
        "If (Fleet Fuel Expenses > EV Charging Net Profit), the DSS triggers an immediate executive alert warning: 'Fleet fuel burn surpasses net EV charging profits. Prioritize fleet EV conversion.' Once EV profits overtake fuel expenditures, the alert switches dynamically to: 'EV charging profits offset X% of internal combustion fleet fuel expenditures.'",
        prefix="CROSS-SECTOR SYNERGY RULE: "
    )
    add_body_p("")
    add_callout(
        "When tyre stock falls below 5 units, the inventory heuristics engine flags an immediate reorder alert detailing brand, pattern, and exact remaining count to prevent stockouts.",
        prefix="INVENTORY REORDER RULE: "
    )

    add_sec_heading("3.3 Tyre Inventory Management & POS Sale Engine", level=2)
    add_body_p(
        "The inventory engine controls the product lifecycle from wholesale procurement to point-of-sale customer settlement:"
    )
    add_body_p("Procurement Ingestion: Captures brand, tread model, tire dimensions (e.g., 205/55R16), unit buying cost, retail price, and starting quantity.", bullet=True)
    add_body_p("Atomic POS Settlement: Initiates a database transaction that acquires a row-lock (SELECT ... FOR UPDATE), validates whether the available quantity meets the requested order, deducts stock atomically, and logs the customer sales voucher.", bullet=True)
    add_body_p("Real-Time Unit Margin Calculation: Computes gross profit per unit dynamically: Profit Margin = Selling Price - Buying Price.", bullet=True)

    add_sec_heading("3.4 Electric Vehicle (EV) Charging Station Operations", level=2)
    add_body_p(
        "Administers multi-standard public charging infrastructure located at the Piliyandala hub:"
    )
    add_body_p("Charging Point Registry: Supports DC Fast CCS2 (120 LKR/kWh), CHAdeMO, and Type 2 AC chargers.", bullet=True)
    add_body_p("Session Logging & Tariff Application: Operators log customer vehicle plate numbers and meter energy readings (kWh). The backend computes: Total Amount = Energy Consumed (kWh) * Tariff Rate.", bullet=True)
    add_body_p("Net Energy Margin Estimation: Calculates electricity tariff margins after factoring Ceylon Electricity Board (CEB) commercial utility tariffs.", bullet=True)

    add_sec_heading("3.5 Commercial Fleet Operations (Uber / PickMe Tracking)", level=2)
    add_body_p(
        "Monitors commercial passenger vehicles operating on digital ride-hailing networks:"
    )
    add_body_p("Vehicle & Driver Registry: Maintains records of registered fleet units (e.g., WP CAB-1234 Toyota Prius) and active driver designations.", bullet=True)
    add_body_p("Trip Gross Earnings Log: Ingests platform daily gross revenues (Uber, PickMe, Private Hire).", bullet=True)
    add_body_p("Operational Cost Vouchers: Tracks itemized operating costs across Fuel, Maintenance, Insurance, and Repairs to accurately calculate net vehicle profit.", bullet=True)

    # -------------------------------------------------------------
    # SECTION 4: DATABASE ARCHITECTURE
    # -------------------------------------------------------------
    add_sec_heading("4. Database Architecture & Relational Schema")
    add_body_p(
        "The persistence layer consists of 8 normalized relational tables in MySQL (InnoDB), strictly enforcing foreign key relationships and data integrity:"
    )

    db_table = doc.add_table(rows=9, cols=3)
    db_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    db_headers = ["Table Name", "Primary Key & Foreign Keys", "Business Function & Data Stored"]
    for j, h in enumerate(db_headers):
        cell = db_table.rows[0].cells[j]
        set_cell_background(cell, "102C57")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    db_rows = [
        ("users", "id (PK)", "System accounts, password hashes, and user roles (Owner, Shop Manager, EV Operator, Driver)."),
        ("tyres", "id (PK)", "Tyre catalog with dimensions, unit buying cost, selling price, and stock quantities."),
        ("tyre_sales", "id (PK), tyre_id (FK -> tyres.id)", "Sales transactions, customer names, quantity sold, and total billed revenue."),
        ("charging_points", "id (PK)", "EV charging bays, connector types (CCS2/AC), base tariff rates, and active states."),
        ("charging_sessions", "id (PK), point_id (FK -> charging_points.id)", "Charging sessions, vehicle license plates, kWh dispensed, and computed billing amounts."),
        ("vehicles", "id (PK)", "Fleet vehicles, license plate numbers, vehicle models, assigned drivers, and status."),
        ("fleet_trips", "id (PK), vehicle_id (FK -> vehicles.id)", "Daily platform earnings logs (Uber, PickMe, Private Hire) with dates and gross revenues."),
        ("fleet_expenses", "id (PK), vehicle_id (FK -> vehicles.id)", "Fleet expenditures categorized by Fuel, Routine Maintenance, Insurance, and Repairs.")
    ]
    for i, row in enumerate(db_rows):
        r_cells = db_table.rows[i+1].cells
        for j, val in enumerate(row):
            cell = r_cells[j]
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            set_cell_background(cell, "F8F9FA" if i % 2 == 0 else "FFFFFF")
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if j == 0:
                r.font.bold = True

    # -------------------------------------------------------------
    # SECTION 5: VERIFICATION & TESTING RESULTS
    # -------------------------------------------------------------
    add_sec_heading("5. Verification, Execution & Testing Results")
    add_body_p(
        "The system has been verified through end-to-end operational test cycles matching actual production figures:"
    )

    test_table = doc.add_table(rows=4, cols=4)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    test_headers = ["Operational Sector", "Gross Revenue (LKR)", "Net Profit (LKR)", "Operational Margin (%)"]
    for j, h in enumerate(test_headers):
        cell = test_table.rows[0].cells[j]
        set_cell_background(cell, "102C57")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    test_rows = [
        ("Tyre Inventory & Sales POS", "Rs. 84,000.00", "Rs. 14,000.00", "16.7%"),
        ("EV Charging Station Infrastructure", "Rs. 4,260.00", "Rs. 1,704.00", "40.0%"),
        ("Fleet Operations (Uber / PickMe)", "Rs. 18,500.00", "Rs. 13,300.00", "71.9%")
    ]
    for i, row in enumerate(test_rows):
        r_cells = test_table.rows[i+1].cells
        for j, val in enumerate(row):
            cell = r_cells[j]
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            set_cell_background(cell, "F8F9FA" if i % 2 == 0 else "FFFFFF")
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            if j == 0:
                r.font.bold = True

    add_body_p(
        "Consolidated Financial Audit Result: Total Business Revenue = Rs. 106,760.00 | Total Net Business Profit = Rs. 29,004.00. The Python engine validated cross-sector synergy coverage at 32.8% of fleet fuel burn covered by EV charging tariffs.",
        bold_prefix="Consolidated Performance: "
    )

    # -------------------------------------------------------------
    # SECTION 6: FUTURE EXPANSIONS & CONCLUSION
    # -------------------------------------------------------------
    add_sec_heading("6. Strategic Recommendations & Conclusion")
    add_body_p(
        "The Integrated Business Decision Support System successfully eliminates cross-sector data fragmentation for commercial multi-unit operations in Piliyandala. Future technical milestones include:"
    )
    add_body_p("Automated IoT Charging Ingestion: Integrating OCPP (Open Charge Point Protocol) to log session metrics directly from EV hardware without manual operator entry.", bullet=True)
    add_body_p("Predictive Maintenance: Using regression models to project fleet tire replacement intervals based on recorded trip mileage and wear rates.", bullet=True)
    add_body_p("Automated End-of-Day Email Dispatch: Leveraging NodeMailer to send daily PDF financial audits directly to stakeholders at 23:59 daily.", bullet=True)

    # Save document
    filename = "Integrated_Business_DSS_Project_Report.docx"
    doc.save(filename)
    print(f"\n[SUCCESS] Professional Word Document created successfully: {os.path.abspath(filename)}")

if __name__ == '__main__':
    create_full_project_report()